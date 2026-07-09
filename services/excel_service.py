from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func

from db.db_manager import get_session
from db.models import AnnualDeclaration, User, UserDeclarationStatus, as_declaration_name

EXCEL_HEADERS = [
    "Staff ID",
    "Name",
    "Email",
    "Department",
    "Status",
    "Has Conflicts",
    "Notify",
    "Remarks",
]

_INVALID_SHEET_CHARS = str.maketrans({c: "_" for c in r'\/?*[]:'})


def _sanitize_sheet_title(title: str) -> str:
    cleaned = title.translate(_INVALID_SHEET_CHARS).strip()
    return (cleaned or "Report")[:31]


def _display_status(entry: UserDeclarationStatus | None) -> str:
    if not entry or entry.status != "completed":
        return "Pending"
    return "Completed"


def _parse_yes_no(value) -> bool:
    if value is None:
        return True
    return str(value).strip().lower() in {"yes", "y", "true", "1"}


def _header_index_map(header_row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is not None:
            mapping[str(cell).strip()] = idx
    return mapping


def parse_excel_counts(file_bytes: bytes) -> tuple[int, int, str]:
    """Return (pending_count, total_count, excel_pending_status) from uploaded file."""
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        return 0, 0, "0/0"

    col_map = _header_index_map(header)
    status_col = col_map.get("Status")

    total = 0
    pending = 0

    for row in rows:
        if not row or not any(row):
            continue
        staff_col = col_map.get("Staff ID")
        if staff_col is not None and not row[staff_col]:
            continue

        total += 1
        if status_col is not None:
            status_val = str(row[status_col] or "").strip().lower()
            if status_val != "completed":
                pending += 1
        else:
            pending += 1

    wb.close()
    return pending, total, f"{pending}/{total}"


async def generate_declaration_report(declaration_id: str) -> BytesIO:
    """Excel template: user metadata + status. Uses write-only mode for large user counts."""

    async with get_session() as session:
        declaration = await session.get(AnnualDeclaration, declaration_id)
        if not declaration:
            raise ValueError("Declaration not found")

        decl_name = as_declaration_name(declaration.declaration_name)
        financial_year = declaration.financial_year

        users = (
            await session.execute(
                select(
                    User.staff_id,
                    User.username,
                    User.email,
                    User.department,
                )
                .where(func.lower(User.status) == "active")
                .order_by(User.staff_id)
            )
        ).all()

        status_rows = (
            await session.execute(
                select(UserDeclarationStatus).where(
                    UserDeclarationStatus.declaration_id == declaration_id
                )
            )
        ).scalars().all()

        status_map = {s.staff_id: s for s in status_rows}

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=_sanitize_sheet_title(f"{decl_name} FY {financial_year}"))
    ws.append(EXCEL_HEADERS)

    for user in users:
        entry = status_map.get(user.staff_id)
        ws.append([
            user.staff_id,
            user.username or "",
            user.email,
            user.department or "",
            _display_status(entry),
            "Yes" if entry and entry.has_conflicts else "No",
            "Yes" if not entry or entry.notify else "No",
            "",
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
